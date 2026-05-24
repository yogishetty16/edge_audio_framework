import os
import sys
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Theme Colors
BLACK = (0, 0, 0)
DARK_BLUE = (31, 73, 125)
GREY = (89, 89, 89)

def clean_xml_string(s):
    """Filter out characters that are not valid in XML."""
    if not isinstance(s, str):
        return s
    return "".join(
        c for c in s
        if ord(c) in (0x9, 0xA, 0xD)
        or (0x20 <= ord(c) <= 0xD7FF)
        or (0xE000 <= ord(c) <= 0xFFFD)
        or (0x10000 <= ord(c) <= 0x10FFFF)
    )

def font(run, size=11, bold=False, italic=False, color=None):
    """Apply styling directly to a run."""
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    else:
        run.font.color.rgb = RGBColor(*BLACK)

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set the internal margins (padding) of a table cell in dxa (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_inline_to_paragraph(p, text, font_size=11):
    """Parse basic markdown inline syntax and append runs to the paragraph."""
    # Split by bold (**), italic (*), inline code (`), and markdown links ([text](url))
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            font(r, font_size, bold=True)
        elif part.startswith('*') and part.endswith('*'):
            r = p.add_run(part[1:-1])
            font(r, font_size, italic=True)
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(font_size - 1.5)
            r.font.color.rgb = RGBColor(*GREY)
        elif part.startswith('[') and ']' in part:
            match = re.match(r'^\[(.*?)\]\((.*?)\)$', part)
            if match:
                anchor, url = match.groups()
                r = p.add_run(anchor)
                font(r, font_size, color=DARK_BLUE)
                r.font.underline = True
            else:
                r = p.add_run(part)
                font(r, font_size)
        else:
            r = p.add_run(part)
            font(r, font_size)

def add_h1(doc, text):
    p = doc.add_heading(level=1)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font(r, 16, bold=True, color=DARK_BLUE)
    return p

def add_h2(doc, text):
    p = doc.add_heading(level=2)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font(r, 14, bold=True, color=DARK_BLUE)
    return p

def add_h3(doc, text):
    p = doc.add_heading(level=3)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font(r, 12, bold=True, color=BLACK)
    return p

def add_blockquote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(6)
    
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            font(r, 10.5, bold=True, italic=True, color=GREY)
        elif part.startswith('*') and part.endswith('*'):
            r = p.add_run(part[1:-1])
            font(r, 10.5, italic=True, color=GREY)
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(*GREY)
        elif part.startswith('[') and ']' in part:
            match = re.match(r'^\[(.*?)\]\((.*?)\)$', part)
            if match:
                anchor, _ = match.groups()
                r = p.add_run(anchor)
                font(r, 10.5, italic=True, color=DARK_BLUE)
                r.font.underline = True
            else:
                r = p.add_run(part)
                font(r, 10.5, italic=True, color=GREY)
        else:
            r = p.add_run(part)
            font(r, 10.5, italic=True, color=GREY)

def add_code_block(doc, lines):
    """Renders a code block inside a shaded 1x1 table cell with custom borders."""
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    
    set_cell_background(cell, "F2F5F8") # Light blue-grey background
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Customize cell borders to a thin grey border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{b_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4') # 0.5 pt width
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'D3D3D3')
        tcBorders.append(border)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(*GREY)

def add_table(doc, table_lines):
    """Parses markdown table lines and adds a styled Word table."""
    rows_data = []
    for line in table_lines:
        if re.match(r"^\|[\s:-|]*\|$", line):
            # Skip separator line
            continue
        cells = [c.strip() for c in line.split("|")]
        if cells and cells[0] == "":
            cells = cells[1:]
        if cells and cells[-1] == "":
            cells = cells[:-1]
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_rows = len(rows_data)
    num_cols = len(rows_data[0])
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Table Grid'
    
    for r_idx, row_data in enumerate(rows_data):
        row = table.rows[r_idx]
        is_header = (r_idx == 0)
        
        for c_idx, val in enumerate(row_data):
            if c_idx >= num_cols:
                break
            cell = row.cells[c_idx]
            
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            
            if is_header:
                set_cell_background(cell, "1F497D") # Theme Dark Blue
            else:
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F2F5F8") # Alternating row shading
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            if is_header:
                # White bold text
                parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', val)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        r = p.add_run(part[2:-2])
                        font(r, 10, bold=True, color=(255, 255, 255))
                    elif part.startswith('*') and part.endswith('*'):
                        r = p.add_run(part[1:-1])
                        font(r, 10, italic=True, color=(255, 255, 255))
                    else:
                        r = p.add_run(part)
                        font(r, 10, bold=True, color=(255, 255, 255))
            else:
                add_inline_to_paragraph(p, val, font_size=10)

def convert_md_to_docx(input_path, output_path):
    print(f"Converting {input_path} to {output_path}...")
    doc = Document()
    
    # Configure page margins
    for sec in doc.sections:
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.17)
        sec.right_margin = Cm(2.54)
        
    # Configure global Normal style defaults
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    
    with open(input_path, "r", encoding="utf-8") as f:
        lines = [clean_xml_string(line) for line in f.readlines()]
        
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i]
        stripped = line.strip()
        
        # 1. Code Block
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue
            
        # 2. Markdown Table
        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue
            
        # 3. Blank line
        if not stripped:
            doc.add_paragraph()
            i += 1
            continue
            
        # 4. Page Break / Horizontal Rule
        if re.match(r"^-{3,}$", stripped):
            doc.add_page_break()
            i += 1
            continue
            
        # 5. Blockquote
        if stripped.startswith(">"):
            add_blockquote(doc, stripped[1:].strip())
            i += 1
            continue
            
        # 6. Headings
        if stripped.startswith("# "):
            add_h1(doc, stripped[2:].strip())
            i += 1
            continue
        elif stripped.startswith("## "):
            add_h2(doc, stripped[3:].strip())
            i += 1
            continue
        elif stripped.startswith("### "):
            add_h3(doc, stripped[4:].strip())
            i += 1
            continue
            
        # 7. Bullet / List Items
        leading_spaces = len(line) - len(line.lstrip(' '))
        indent_level = leading_spaces // 2
        
        if stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            if indent_level > 0:
                p.paragraph_format.left_indent = Cm(0.5 + 0.5 * indent_level)
            add_inline_to_paragraph(p, stripped[2:].strip())
            i += 1
            continue
            
        # 8. Numbered list
        match_num = re.match(r"^(\d+)\.\s(.*)$", stripped)
        if match_num:
            num, content = match_num.groups()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            if indent_level > 0:
                p.paragraph_format.left_indent = Cm(0.5 + 0.5 * indent_level)
            add_inline_to_paragraph(p, content.strip())
            i += 1
            continue
            
        # 9. Normal paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        add_inline_to_paragraph(p, stripped)
        i += 1
        
    doc.save(output_path)
    print(f"Successfully saved {output_path}!")

if __name__ == "__main__":
    # If arguments are passed, convert specified files
    if len(sys.argv) >= 3:
        convert_md_to_docx(sys.argv[1], sys.argv[2])
    else:
        # Default: Convert README and DOCUMENTATION
        convert_md_to_docx("README.md", "README.docx")
        convert_md_to_docx("DOCUMENTATION.md", "DOCUMENTATION.docx")
